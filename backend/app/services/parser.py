import re
from pathlib import Path

from loguru import logger

from app.core.config import get_settings
from app.utils.file_helper import get_file_extension

settings = get_settings()


def _normalize_text(text: str) -> str:
    """Collapse whitespace and remove noisy empty lines."""
    lines = []
    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def parse_pdf(file_path: str) -> str:
    text_parts: list[str] = []

    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        logger.info(f"pdfplumber extracted text from {file_path}")
    except Exception as e:
        logger.warning(f"pdfplumber failed on {file_path}: {e}")

    if not text_parts:
        try:
            import fitz

            doc = fitz.open(file_path)
            for page in doc:
                page_text = page.get_text()
                if page_text:
                    text_parts.append(page_text)
            doc.close()
            logger.info(f"PyMuPDF extracted text from {file_path}")
        except Exception as e:
            logger.error(f"PyMuPDF also failed on {file_path}: {e}")

    return _normalize_text("\n".join(text_parts))


def parse_docx(file_path: str) -> str:
    text_parts: list[str] = []

    try:
        from docx import Document

        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    text_parts.append(" | ".join(cells))

        logger.info(f"python-docx extracted text from {file_path}")
    except Exception as e:
        logger.error(f"python-docx failed on {file_path}: {e}")

    return _normalize_text("\n".join(text_parts))


def parse_resume(file_path: str) -> str:
    ext = get_file_extension(file_path)
    logger.info(f"Parsing resume: {file_path} (type: {ext})")

    if ext == "pdf":
        text = parse_pdf(file_path)
    elif ext == "docx":
        text = parse_docx(file_path)
    else:
        logger.error(f"Unsupported file type: {ext}")
        raise ValueError(f"Unsupported file type: {ext}")

    if not text.strip():
        logger.warning(f"No text extracted from {file_path}")
        raise ValueError(
            "Could not extract any text from the resume file. "
            "The file may be a scanned image; please upload a text-based PDF or DOCX."
        )

    if len(text) > settings.MAX_RESUME_TEXT_CHARS:
        logger.warning(
            f"Truncating resume text ({len(text)} chars) to {settings.MAX_RESUME_TEXT_CHARS}"
        )
        text = text[: settings.MAX_RESUME_TEXT_CHARS]

    logger.info(f"Successfully parsed {file_path}: {len(text)} characters extracted")
    return text
